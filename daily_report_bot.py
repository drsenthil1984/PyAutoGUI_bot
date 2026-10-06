import pyautogui
import pyperclip
import time
from datetime import datetime

from openpyxl import Workbook
import re

pyautogui.FAILSAFE = True
pyautogui.PAUSE = 0.5

# Current date/time
now = datetime.now()

current_datetime = now.strftime("%Y-%m-%d %H:%M:%S")
current_date = now.strftime("%Y-%m-%d")

excel_filename = f"daily_report_{current_date}.xlsx"
screenshot_filename = f"daily_report_{current_date}.png"

comment = "Screen Stock Data"

print("step 1:open the chrome browser...")
time.sleep(2)   
pyautogui.hotkey('command', 'space', interval=0.1)
time.sleep(1)

pyautogui.write('chrome', interval=0.15)
pyautogui.press('enter')
time.sleep(3)

print("step 2:open the website...")
pyautogui.hotkey('command', 't', interval=0.1)
time.sleep(1)
pyautogui.write('https://www.screener.in/market/IN07/IN0702/IN070201/', interval=0.15)
pyautogui.press('enter')

time.sleep(6)


print("Step 3: Copy website data...")

# Copy complete webpage
pyautogui.hotkey('command', 'a')
time.sleep(1)

pyautogui.hotkey('command', 'c')
time.sleep(2)

page_data = pyperclip.paste()

print("Website data copied successfully")


# Convert webpage into clean lines
lines = [
    line.strip()
    for line in page_data.splitlines()
    if line.strip()
]


# Find Screener table header
start_index = None

for i, line in enumerate(lines):
    if "S.No." in line and "Company" in line and "CMP" in line:
        start_index = i
        break


if start_index is None:
    print("ERROR: Stock table not found")
    stock_rows = []

else:
    print("Stock table found")

    # Capture next 20 rows
    stock_rows = lines[start_index + 1:start_index + 21]

    print("Rows captured:")
    for row in stock_rows:
        print(row)


print("Step 4: Create Excel report...")

# Create workbook
wb = Workbook()
ws = wb.active
ws.title = "Daily Report"


# Excel headers
headers = [
    "Date & Time",
    "S.No.",
    "Company",
    "CMP Rs.",
    "P/E",
    "Mar Cap Rs.Cr.",
    "Div Yld %",
    "NP Qtr Rs.Cr.",
    "Qtr Profit Var %",
    "Sales Qtr Rs.Cr.",
    "Qtr Sales Var %",
    "ROCE %",
    "Comment"
]


# Write headers
for col_num, header in enumerate(headers, start=1):
    ws.cell(row=1, column=col_num, value=header)


print("Step 5: Write stock data...")

excel_row = 2

for row in stock_rows:

    # Split based on tabs OR multiple spaces
    columns = re.split(r'\t+|\s{2,}', row.strip())

    print("Parsed row:", columns)

    # Date/time
    ws.cell(row=excel_row, column=1, value=current_datetime)

    # Write Screener values
    for col_num, value in enumerate(columns, start=2):
        ws.cell(row=excel_row, column=col_num, value=value)

    # Comment
    ws.cell(row=excel_row, column=13, value=comment)

    excel_row += 1


# Adjust column width
for column_cells in ws.columns:

    max_length = 0
    column_letter = column_cells[0].column_letter

    for cell in column_cells:

        if cell.value is not None:
            max_length = max(
                max_length,
                len(str(cell.value))
            )

    ws.column_dimensions[column_letter].width = min(
        max_length + 2,
        40
    )


# Save as Excel file
excel_filename = f"daily_report_{current_date}.xlsx"

wb.save(excel_filename)

print("--------------------------------")
print("Excel file created successfully:")
print(excel_filename)
print("--------------------------------")